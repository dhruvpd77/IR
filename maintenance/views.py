import csv
from datetime import datetime, timedelta
from io import TextIOWrapper

from django.contrib import messages
from django.core.paginator import Paginator
from django.contrib.auth import login, logout, update_session_auth_hash
from django.contrib.auth.decorators import login_required
from django.contrib.auth.forms import AuthenticationForm
from django.db.models import Case, Count, F, IntegerField, Prefetch, Q, Value, When
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from .decorators import roles_required
from .forms import InventoryForm, PasswordForm, UserProvisionForm
from .models import Asset, Category, InventoryItem, PreventiveTask, StockMovement, SubCategory, Ticket, TicketMedia, User, format_duration
from .services import (
    _cmms_triples,
    add_months,
    apply_stock_change,
    import_catalog,
    import_cmms_triples,
    import_pm_rows,
    next_ticket_no,
    schedule_month_label,
    validate_media,
)

OPEN_STATUSES = [Ticket.Status.PENDING, Ticket.Status.ATTEMPTED]


def _safe_next(request):
    target = request.POST.get('next') or request.GET.get('next') or ''
    if target and url_has_allowed_host_and_scheme(target, allowed_hosts={request.get_host()}):
        return target
    return ''


def login_view(request):
    if request.user.is_authenticated:
        return redirect('home')
    form = AuthenticationForm(request, data=request.POST or None)
    if request.method == 'POST' and form.is_valid():
        login(request, form.get_user())
        return redirect(_safe_next(request) or 'home')
    return render(request, 'auth/login.html', {'form': form})


@require_POST
def logout_view(request):
    logout(request)
    return redirect('login')


def _scoped_tickets(user):
    qs = Ticket.objects.select_related('asset', 'category', 'subcategory', 'raised_by')
    if user.role == User.Role.EMPLOYEE:
        qs = qs.filter(raised_by=user)
    return qs


@login_required
def home(request):
    tickets = _scoped_tickets(request.user)
    grouped = {row['status']: row['c'] for row in tickets.values('status').annotate(c=Count('id'))}
    for key in (Ticket.Status.PENDING, Ticket.Status.ATTEMPTED, Ticket.Status.COMPLETED):
        grouped.setdefault(key, 0)
    start = timezone.localtime().replace(hour=0, minute=0, second=0, microsecond=0)
    done = [t for t in tickets.filter(status=Ticket.Status.COMPLETED, completed_at__isnull=False) if t.attempted_at]
    if done:
        avg_seconds = sum((t.completed_at - t.attempted_at).total_seconds() for t in done) / len(done)
        avg_idle = format_duration(timedelta(seconds=avg_seconds))
    else:
        avg_idle = '—'
    rows = list(tickets.values('category_id', 'category__name').annotate(c=Count('id')).order_by('-c'))
    top = max((row['c'] for row in rows), default=1) or 1
    bars = [
        {'id': row['category_id'], 'name': row['category__name'], 'count': row['c'], 'pct': round(row['c'] / top * 100)}
        for row in rows
    ]
    low_items = []
    people = None
    if request.user.is_admin_role:
        low_items = list(
            InventoryItem.objects.filter(is_active=True, quantity__lte=F('min_quantity')).order_by('quantity')[:6]
        )
    if request.user.is_superadmin:
        people = {
            'admin': User.objects.filter(role=User.Role.ADMIN, is_active=True).count(),
            'employee': User.objects.filter(role=User.Role.EMPLOYEE, is_active=True).count(),
        }
    context = {
        'counts': grouped,
        'machines_down': tickets.filter(status=Ticket.Status.ATTEMPTED).values('asset_id').distinct().count(),
        'completed_today': tickets.filter(status=Ticket.Status.COMPLETED, completed_at__gte=start).count(),
        'avg_idle': avg_idle,
        'open_tickets': tickets.filter(status__in=OPEN_STATUSES).order_by('created_at')[:8],
        'recent_completed': tickets.filter(status=Ticket.Status.COMPLETED).order_by('-completed_at')[:4],
        'bars': bars,
        'low_items': low_items,
        'people': people,
    }
    return render(request, 'dashboard/home.html', context)


@login_required
def ticket_list(request):
    status = request.GET.get('status', 'all')
    if status not in {'all', 'pending', 'attempted', 'completed', 'rejected'}:
        status = 'all'
    q = request.GET.get('q', '').strip()
    priority = request.GET.get('priority', '')
    category_id = request.GET.get('category', '')
    base = _scoped_tickets(request.user)
    counts = {row['status']: row['c'] for row in base.values('status').annotate(c=Count('id'))}
    for key in ('pending', 'attempted', 'completed', 'rejected'):
        counts.setdefault(key, 0)
    counts['all'] = counts['pending'] + counts['attempted'] + counts['completed'] + counts['rejected']
    tickets = base
    if status != 'all':
        tickets = tickets.filter(status=status)
    if priority in dict(Ticket.Priority.choices):
        tickets = tickets.filter(priority=priority)
    if category_id.isdigit():
        tickets = tickets.filter(category_id=int(category_id))
    if q:
        tickets = tickets.filter(
            Q(ticket_no__icontains=q)
            | Q(asset__code__icontains=q)
            | Q(asset__name__icontains=q)
            | Q(asset__location__icontains=q)
            | Q(remarks__icontains=q)
            | Q(category__name__icontains=q)
            | Q(subcategory__name__icontains=q)
            | Q(raised_by__first_name__icontains=q)
            | Q(raised_by__last_name__icontains=q)
        )
    tickets = tickets.annotate(
        status_order=Case(
            When(status=Ticket.Status.PENDING, then=Value(0)),
            When(status=Ticket.Status.ATTEMPTED, then=Value(1)),
            When(status=Ticket.Status.REJECTED, then=Value(2)),
            When(status=Ticket.Status.COMPLETED, then=Value(3)),
            default=Value(4),
            output_field=IntegerField(),
        )
    ).order_by('status_order', '-created_at')
    return render(request, 'tickets/list.html', {
        'tickets': tickets,
        'counts': counts,
        'status': status,
        'q': q,
        'priority': priority,
        'category_id': category_id,
        'categories': Category.objects.filter(is_active=True),
        'priorities': Ticket.Priority.choices,
        'shown': tickets.count(),
    })


@login_required
def ticket_create(request):
    categories = Category.objects.filter(is_active=True).order_by('sort_order', 'name')
    selected = {
        'category': '',
        'subcategory': '',
        'asset': '',
        'remarks': '',
        'priority': Ticket.Priority.MEDIUM,
    }
    names = {'cat': '', 'sub': '', 'asset': ''}
    errors = []
    asset_open = False
    if request.method == 'POST':
        selected = {
            'category': request.POST.get('category', ''),
            'subcategory': request.POST.get('subcategory', ''),
            'asset': request.POST.get('asset', ''),
            'remarks': (request.POST.get('remarks') or '').strip(),
            'priority': request.POST.get('priority') or Ticket.Priority.MEDIUM,
        }
        category = subcategory = asset = None
        if selected['category'].isdigit():
            category = Category.objects.filter(pk=int(selected['category']), is_active=True).first()
        if category and selected['subcategory'].isdigit():
            subcategory = SubCategory.objects.filter(
                pk=int(selected['subcategory']), category=category, is_active=True
            ).first()
        if subcategory and selected['asset'].isdigit():
            asset = Asset.objects.filter(pk=int(selected['asset']), subcategory=subcategory, is_active=True).first()
        if category:
            names['cat'] = category.name
        if subcategory:
            names['sub'] = subcategory.name
        if asset:
            names['asset'] = f'{asset.code} · {asset.name}'
            asset_open = Ticket.objects.filter(asset=asset, status__in=OPEN_STATUSES).exists()
        if not (category and subcategory and asset):
            errors.append('Select a value stream, an equipment name, and a location.')
        if len(selected['remarks']) < 8:
            errors.append('Write a remark so maintenance knows what you saw.')
        if selected['priority'] not in dict(Ticket.Priority.choices):
            selected['priority'] = Ticket.Priority.MEDIUM
        files = request.FILES.getlist('media')
        if len(files) > 8:
            errors.append('You can attach up to 8 photos or videos.')
        for uploaded in files:
            problem = validate_media(uploaded)
            if problem:
                errors.append(problem)
        if not errors:
            ticket = Ticket.objects.create(
                ticket_no=next_ticket_no(),
                raised_by=request.user,
                asset=asset,
                category=category,
                subcategory=subcategory,
                priority=selected['priority'],
                remarks=selected['remarks'],
            )
            for uploaded in files:
                TicketMedia.objects.create(ticket=ticket, file=uploaded, uploaded_by=request.user)
            messages.success(request, f'{ticket.ticket_no} was received. Idle time starts when an admin accepts it.')
            return redirect(ticket.get_absolute_url())
    resume_step = 1
    if selected['category'] and selected['subcategory'] and selected['asset']:
        resume_step = 4
    elif selected['category'] and selected['subcategory']:
        resume_step = 3
    elif selected['category']:
        resume_step = 2
    return render(request, 'tickets/form.html', {
        'categories': categories,
        'selected': selected,
        'names': names,
        'errors': errors,
        'resume_step': resume_step if errors else 1,
        'asset_open': asset_open,
        'priorities': Ticket.Priority.choices,
    })


def _get_ticket_for_user(request, pk):
    ticket = get_object_or_404(
        Ticket.objects.select_related(
            'asset', 'category', 'subcategory', 'raised_by', 'attempted_by', 'completed_by', 'rejected_by'
        ).prefetch_related(
            'media',
            Prefetch(
                'movements',
                queryset=StockMovement.objects.select_related('item', 'created_by').order_by('created_at'),
            ),
        ),
        pk=pk,
    )
    if request.user.role == User.Role.EMPLOYEE and ticket.raised_by_id != request.user.id:
        return None
    return ticket


@login_required
def ticket_detail(request, pk):
    ticket = _get_ticket_for_user(request, pk)
    if ticket is None:
        return render(request, 'errors/403.html', status=403)
    parts = InventoryItem.objects.filter(is_active=True).order_by('name') if request.user.is_admin_role else []
    return render(request, 'tickets/detail.html', {'ticket': ticket, 'parts': parts})


@roles_required(User.Role.SUPERADMIN, User.Role.ADMIN)
@require_POST
def ticket_attempt(request, pk):
    ticket = get_object_or_404(Ticket, pk=pk)
    if ticket.status != Ticket.Status.PENDING:
        messages.warning(request, 'This request is no longer waiting to be accepted.')
        return redirect(ticket.get_absolute_url())
    ticket.status = Ticket.Status.ATTEMPTED
    ticket.attempted_at = timezone.now()
    ticket.attempted_by = request.user
    note = (request.POST.get('attempt_notes') or '').strip()
    if note:
        ticket.attempt_notes = note
    ticket.save()
    messages.success(request, f'{ticket.ticket_no} accepted. Idle time has started. Add notes and parts, then mark it done.')
    return redirect(ticket.get_absolute_url())


@roles_required(User.Role.SUPERADMIN, User.Role.ADMIN)
@require_POST
def ticket_reject(request, pk):
    ticket = get_object_or_404(Ticket, pk=pk)
    if ticket.status != Ticket.Status.PENDING:
        messages.warning(request, 'Only a new request can be rejected.')
        return redirect(ticket.get_absolute_url())
    reason = (request.POST.get('reject_reason') or '').strip()
    if len(reason) < 5:
        messages.error(request, 'Add a short reason before rejecting the request.')
        return redirect(ticket.get_absolute_url())
    ticket.status = Ticket.Status.REJECTED
    ticket.rejected_at = timezone.now()
    ticket.rejected_by = request.user
    ticket.reject_reason = reason
    ticket.save()
    messages.success(request, f'{ticket.ticket_no} was rejected.')
    return redirect('ticket_list')


@roles_required(User.Role.SUPERADMIN, User.Role.ADMIN)
@require_POST
def ticket_notes(request, pk):
    ticket = get_object_or_404(Ticket, pk=pk)
    if ticket.status != Ticket.Status.ATTEMPTED:
        messages.error(request, 'Accept the request before adding work notes.')
        return redirect(ticket.get_absolute_url())
    ticket.attempt_notes = (request.POST.get('attempt_notes') or '').strip()
    ticket.save(update_fields=['attempt_notes'])
    messages.success(request, 'Work notes saved.')
    return redirect(ticket.get_absolute_url())


@roles_required(User.Role.SUPERADMIN, User.Role.ADMIN)
@require_POST
def ticket_complete(request, pk):
    ticket = get_object_or_404(Ticket, pk=pk)
    if ticket.status != Ticket.Status.ATTEMPTED:
        messages.error(request, 'Accept the request before marking it done.')
        return redirect(ticket.get_absolute_url())
    notes = (request.POST.get('completion_notes') or '').strip()
    if len(notes) < 5:
        messages.error(request, 'Add a short completion note before closing the job.')
        return redirect(ticket.get_absolute_url())
    ticket.status = Ticket.Status.COMPLETED
    ticket.completed_at = timezone.now()
    ticket.completed_by = request.user
    ticket.completion_notes = notes
    ticket.save()
    messages.success(request, f'{ticket.ticket_no} is completed. Idle time is frozen at {ticket.idle_display}.')
    return redirect(ticket.get_absolute_url())


@roles_required(User.Role.SUPERADMIN, User.Role.ADMIN)
@require_POST
def ticket_add_part(request, pk):
    ticket = get_object_or_404(Ticket, pk=pk)
    if ticket.status != Ticket.Status.ATTEMPTED:
        messages.error(request, 'Accept the request before issuing parts.')
        return redirect(ticket.get_absolute_url())
    item = InventoryItem.objects.filter(pk=request.POST.get('item'), is_active=True).first()
    if item is None:
        messages.error(request, 'Choose a part from inventory.')
        return redirect(ticket.get_absolute_url())
    try:
        qty = int(request.POST.get('quantity', '0'))
    except ValueError:
        qty = 0
    if qty < 1:
        messages.error(request, 'Enter a quantity of at least 1.')
        return redirect(ticket.get_absolute_url())
    try:
        apply_stock_change(
            item=item,
            quantity_change=-qty,
            reason=StockMovement.Reason.JOB,
            user=request.user,
            ticket=ticket,
            note=(request.POST.get('note') or '').strip(),
        )
    except ValueError as exc:
        messages.error(request, str(exc))
    else:
        messages.success(request, f'Issued {qty} × {item.name}. On-hand quantity is updated.')
    return redirect(ticket.get_absolute_url())


@roles_required(User.Role.SUPERADMIN, User.Role.ADMIN)
def inventory_list(request):
    q = request.GET.get('q', '').strip()
    low_only = request.GET.get('low') == '1'
    items = InventoryItem.objects.filter(is_active=True)
    if q:
        items = items.filter(Q(sku__icontains=q) | Q(name__icontains=q) | Q(location__icontains=q))
    if low_only:
        items = items.filter(quantity__lte=F('min_quantity'))
    low_count = InventoryItem.objects.filter(is_active=True, quantity__lte=F('min_quantity')).count()
    return render(request, 'inventory/list.html', {
        'items': items,
        'q': q,
        'low_only': low_only,
        'low_count': low_count,
    })


@roles_required(User.Role.SUPERADMIN, User.Role.ADMIN)
def inventory_create(request):
    form = InventoryForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        item = form.save()
        if item.quantity:
            StockMovement.objects.create(
                item=item,
                quantity_change=item.quantity,
                balance_after=item.quantity,
                reason=StockMovement.Reason.RECEIPT,
                note='Opening stock',
                created_by=request.user,
            )
        messages.success(request, f'{item.name} is in inventory with {item.quantity} {item.unit}.')
        return redirect('inventory_detail', pk=item.pk)
    return render(request, 'inventory/form.html', {'form': form})


@roles_required(User.Role.SUPERADMIN, User.Role.ADMIN)
def inventory_detail(request, pk):
    item = get_object_or_404(InventoryItem, pk=pk)
    if request.method == 'POST':
        direction = request.POST.get('direction')
        note = (request.POST.get('note') or '').strip()
        try:
            qty = int(request.POST.get('quantity', '0'))
        except ValueError:
            qty = 0
        if qty < 1:
            messages.error(request, 'Enter a quantity of at least 1.')
        elif direction not in {'receive', 'issue'}:
            messages.error(request, 'Choose receive or remove.')
        else:
            change = qty if direction == 'receive' else -qty
            reason = StockMovement.Reason.RECEIPT if direction == 'receive' else StockMovement.Reason.ADJUST
            try:
                apply_stock_change(
                    item=item,
                    quantity_change=change,
                    reason=reason,
                    user=request.user,
                    note=note,
                )
            except ValueError as exc:
                messages.error(request, str(exc))
            else:
                messages.success(request, 'Stock quantity updated.')
        return redirect('inventory_detail', pk=item.pk)
    movements = item.movements.select_related('created_by', 'ticket')
    return render(request, 'inventory/detail.html', {'item': item, 'movements': movements})


@roles_required(User.Role.SUPERADMIN)
def user_list(request):
    q = request.GET.get('q', '').strip()
    people = User.objects.annotate(
        role_order=Case(
            When(role=User.Role.SUPERADMIN, then=Value(0)),
            When(role=User.Role.ADMIN, then=Value(1)),
            default=Value(2),
            output_field=IntegerField(),
        )
    ).order_by('role_order', 'first_name', 'username')
    if q:
        people = people.filter(
            Q(username__icontains=q)
            | Q(first_name__icontains=q)
            | Q(last_name__icontains=q)
            | Q(employee_code__icontains=q)
            | Q(department__icontains=q)
        )
    return render(request, 'users/list.html', {'people': people, 'q': q})


@roles_required(User.Role.SUPERADMIN)
def user_create(request):
    form = UserProvisionForm(request.POST or None, actor=request.user)
    if request.method == 'POST' and form.is_valid():
        person = form.save()
        messages.success(
            request,
            f'{person.get_full_name()} can sign in as {person.username}. Share the ID and password with them.',
        )
        return redirect('user_list')
    return render(request, 'users/form.html', {'form': form})


@roles_required(User.Role.SUPERADMIN)
def user_password(request, pk):
    person = get_object_or_404(User, pk=pk)
    form = PasswordForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        person.set_password(form.cleaned_data['password'])
        person.save()
        if person.pk == request.user.pk:
            update_session_auth_hash(request, person)
        messages.success(request, f'Password updated for {person.username}.')
        return redirect('user_list')
    return render(request, 'users/password.html', {'form': form, 'person': person})


@roles_required(User.Role.SUPERADMIN)
@require_POST
def user_toggle(request, pk):
    person = get_object_or_404(User, pk=pk)
    if person.pk == request.user.pk:
        messages.error(request, 'You cannot deactivate your own sign-in.')
        return redirect('user_list')
    if person.role == User.Role.SUPERADMIN and person.is_active:
        if User.objects.filter(role=User.Role.SUPERADMIN, is_active=True).count() <= 1:
            messages.error(request, 'Keep at least one active super admin.')
            return redirect('user_list')
    person.is_active = not person.is_active
    person.save(update_fields=['is_active'])
    state = 'activated' if person.is_active else 'deactivated'
    messages.success(request, f'{person.username} is {state}.')
    return redirect('user_list')


def _board_context(request):
    status = request.GET.get('status', 'open')
    if status not in {'open', 'pending', 'attempted', 'completed'}:
        status = 'open'
    base = Ticket.objects.select_related('asset', 'category', 'subcategory', 'raised_by')
    start = timezone.localtime().replace(hour=0, minute=0, second=0, microsecond=0)
    open_qs = base.filter(status__in=OPEN_STATUSES)
    counts = {
        'pending': open_qs.filter(status=Ticket.Status.PENDING).count(),
        'attempted': open_qs.filter(status=Ticket.Status.ATTEMPTED).count(),
        'machines': open_qs.filter(status=Ticket.Status.ATTEMPTED).values('asset_id').distinct().count(),
        'completed_today': base.filter(status=Ticket.Status.COMPLETED, completed_at__gte=start).count(),
        'low_stock': InventoryItem.objects.filter(is_active=True, quantity__lte=F('min_quantity')).count(),
    }
    if status == 'pending':
        rows = open_qs.filter(status=Ticket.Status.PENDING).order_by('created_at')
    elif status == 'attempted':
        rows = open_qs.filter(status=Ticket.Status.ATTEMPTED).order_by('created_at')
    elif status == 'completed':
        rows = base.filter(status=Ticket.Status.COMPLETED).order_by('-completed_at')[:12]
    else:
        rows = open_qs.order_by('created_at')
    return {'rows': rows, 'counts': counts, 'board_status': status}


@roles_required(User.Role.SUPERADMIN, User.Role.ADMIN)
def tv_board(request):
    context = _board_context(request)
    template = 'tv/_board_body.html' if request.GET.get('partial') == '1' else 'tv/board.html'
    response = render(request, template, context)
    response['Cache-Control'] = 'no-store'
    return response


def _import_catalog_rows(rows):
    for row in rows:
        labels = [str(cell or '').strip().lower() for cell in (row or [])]
        if 'equipment name' in labels and any('valuestream' in label or label == 'value stream' for label in labels):
            return import_cmms_triples(_cmms_triples(rows))
    if not rows:
        raise ValueError('The file is empty.')
    return import_catalog(rows[0], rows[1:])


@roles_required(User.Role.SUPERADMIN)
def import_categories(request):
    categories = Category.objects.prefetch_related('subcategories__assets').order_by('sort_order', 'name')
    if request.method == 'POST':
        upload = request.FILES.get('file')
        if not upload:
            messages.error(request, 'Choose an Excel or CSV file.')
            return redirect('import_categories')
        name = upload.name.lower()
        try:
            if name.endswith('.csv'):
                reader = csv.reader(TextIOWrapper(upload.file, encoding='utf-8-sig'))
                rows = list(reader)
                if not rows:
                    raise ValueError('The file is empty.')
                result = _import_catalog_rows(rows)
            elif name.endswith('.xlsx'):
                from openpyxl import load_workbook
                wb = load_workbook(upload, read_only=True, data_only=True)
                rows = list(wb.active.iter_rows(values_only=True))
                wb.close()
                if not rows:
                    raise ValueError('The file is empty.')
                result = _import_catalog_rows(rows)
            else:
                raise ValueError('Upload a .xlsx or .csv file. In Excel, use Save As if the file is .xls.')
        except (ValueError, InvalidFileException, UnicodeDecodeError, csv.Error) as exc:
            messages.error(request, str(exc) or 'That file could not be read.')
            return redirect('import_categories')
        if 'locations' in result:
            messages.success(
                request,
                f"Added {result['categories']} value streams, {result['subcategories']} equipment names, "
                f"and {result['locations']} locations. Skipped {result['skipped']} rows already in the catalog.",
            )
        else:
            messages.success(
                request,
                f"Added {result['categories']} categories and {result['subcategories']} subcategories. "
                f"Skipped {result['skipped']} rows already in the catalog.",
            )
        return redirect('import_categories')
    return render(request, 'catalog/import.html', {'categories': categories})


@roles_required(User.Role.SUPERADMIN)
def download_category_template(request):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Categories'
    ws.append(['ValueStream', 'Equipment Name', 'Location'])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='0B1F3A')
    for row in (
        ('T 30', 'Forklift - 17 (Toyota)', 'T30 Assembly'),
        ('Rotary', 'Dock Leveler', 'Rotary Store'),
        ('Utility', 'MS sliding gate', 'Admin Building'),
    ):
        ws.append(row)
    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 28
    ws.column_dimensions['C'].width = 62
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="IR_category_template.xlsx"'
    wb.save(response)
    return response


@login_required
def api_subcategories(request, pk):
    subs = SubCategory.objects.filter(category_id=pk, is_active=True).order_by('sort_order', 'name')
    return JsonResponse({'results': [
        {'id': row.id, 'name': row.name, 'description': row.description}
        for row in subs
    ]})


@login_required
def api_assets(request, pk):
    assets = list(Asset.objects.filter(subcategory_id=pk, is_active=True).order_by('location', 'code'))
    open_map = {}
    if assets:
        for asset_id, ticket_no in Ticket.objects.filter(asset__in=assets, status__in=OPEN_STATUSES).values_list('asset_id', 'ticket_no'):
            open_map.setdefault(asset_id, ticket_no)
    return JsonResponse({'results': [
        {
            'id': asset.id,
            'code': asset.code,
            'name': asset.name,
            'location': asset.location,
            'open': asset.id in open_map,
            'open_ticket': open_map.get(asset.id, ''),
        }
        for asset in assets
    ]})


def _month_sort_key(label):
    names = {'Jan': 1, 'Feb': 2, 'Mar': 3, 'Apr': 4, 'May': 5, 'Jun': 6, 'Jul': 7, 'Aug': 8, 'Sep': 9, 'Oct': 10, 'Nov': 11, 'Dec': 12}
    text = (label or '').strip()
    year = 99
    if "'" in text:
        tail = text.split("'")[-1][:2]
        if tail.isdigit():
            year = int(tail)
    return (year, names.get(text[:3], 99), text)


def _pct(part, total):
    if not total:
        return 0.0
    return round(part * 100 / total, 1)


def _month_span(label):
    names = {'Jan': 1, 'Feb': 2, 'Mar': 3, 'Apr': 4, 'May': 5, 'Jun': 6, 'Jul': 7, 'Aug': 8, 'Sep': 9, 'Oct': 10, 'Nov': 11, 'Dec': 12}
    text = (label or '').strip()
    month_no = names.get(text[:3])
    if not month_no or "'" not in text:
        return None
    tail = text.split("'")[-1][:2]
    if not tail.isdigit():
        return None
    year = 2000 + int(tail)
    start = datetime(year, month_no, 1).date()
    if month_no == 12:
        end = datetime(year, 12, 31).date()
    else:
        end = (datetime(year, month_no + 1, 1) - timedelta(days=1)).date()
    return start, end


def _week_choices(qs, span=None):
    if span:
        qs = qs.filter(plan_date__gte=span[0], plan_date__lte=span[1])
    found = {}
    for plan in qs.order_by('plan_date').values_list('plan_date', flat=True).distinct():
        iso = plan.isocalendar()
        key = f'{iso.year}-W{iso.week:02d}'
        if key in found:
            continue
        monday = plan - timedelta(days=plan.weekday())
        sunday = monday + timedelta(days=6)
        show_start = max(monday, span[0]) if span else monday
        show_end = min(sunday, span[1]) if span else sunday
        found[key] = {
            'key': key,
            'label': f'Week {iso.week} · {show_start:%d %b %Y} – {show_end:%d %b %Y}',
            'start': monday,
        }
    return sorted(found.values(), key=lambda item: item['start'])


def _apply_period(qs, month, week, stream='', vendor=''):
    months = sorted({name for name in qs.values_list('schedule_month', flat=True) if name}, key=_month_sort_key)
    if month not in months:
        month = ''
    scoped = qs.filter(schedule_month=month) if month else qs
    span = _month_span(month) if month else None
    weeks = _week_choices(scoped, span)
    if week not in {item['key'] for item in weeks}:
        week = ''
    if week:
        year_text, week_text = week.split('-W')
        scoped = scoped.filter(plan_date__iso_year=int(year_text), plan_date__week=int(week_text))
        if span:
            scoped = scoped.filter(plan_date__gte=span[0], plan_date__lte=span[1])
    stream_names = sorted({name for name in scoped.values_list('value_stream', flat=True) if name})
    if stream not in stream_names:
        stream = ''
    if stream:
        scoped = scoped.filter(value_stream=stream)
    vendor_names = sorted({
        name for name in scoped.filter(work_type__istartswith='external').exclude(vendor_name='').values_list('vendor_name', flat=True)
    })
    if vendor not in vendor_names:
        vendor = ''
    if vendor:
        scoped = scoped.filter(work_type__istartswith='external', vendor_name=vendor)
    label = next((item['label'] for item in weeks if item['key'] == week), '')
    return scoped, months, month, weeks, week, label, stream_names, stream, vendor_names, vendor


def _pm_base():
    today = timezone.localdate()
    soon = today + timedelta(days=7)
    base = PreventiveTask.objects.all()
    return base, today, soon


def _pm_analytics(base, today, soon):
    overdue_q = Q(completed_on__isnull=True, plan_date__lt=today)
    due_q = Q(completed_on__isnull=True, plan_date__gte=today, plan_date__lte=soon)
    upcoming_q = Q(completed_on__isnull=True, plan_date__gt=soon)
    done_q = Q(completed_on__isnull=False)
    counts = base.aggregate(
        overdue=Count('id', filter=overdue_q),
        due=Count('id', filter=due_q),
        upcoming=Count('id', filter=upcoming_q),
        completed=Count('id', filter=done_q),
        month=Count('id', filter=Q(completed_on__year=today.year, completed_on__month=today.month)),
    )
    counts['all'] = counts['overdue'] + counts['due'] + counts['upcoming'] + counts['completed']
    counts['rate'] = _pct(counts['completed'], counts['all'])
    counts['overdue_pct'] = _pct(counts['overdue'], counts['all'])
    def grouped(qs, field):
        rows = []
        for row in qs.values(field).annotate(
            overdue=Count('id', filter=overdue_q),
            due=Count('id', filter=due_q),
            upcoming=Count('id', filter=upcoming_q),
            completed=Count('id', filter=done_q),
            total=Count('id'),
        ).order_by('-overdue', '-completed', field):
            total = row['total'] or 0
            row['done_pct'] = _pct(row['completed'], total)
            row['overdue_pct'] = _pct(row['overdue'], total)
            rows.append(row)
        return rows

    streams = grouped(base, 'value_stream')
    vendors = grouped(
        base.filter(work_type__istartswith='external').exclude(vendor_name=''),
        'vendor_name',
    )
    return counts, streams, vendors, overdue_q, due_q, upcoming_q, done_q


@roles_required(User.Role.SUPERADMIN, User.Role.ADMIN)
def pm_list(request):
    base, today, soon = _pm_base()
    view = request.GET.get('view', 'due')
    if view not in {'overdue', 'due', 'reminder', 'upcoming', 'completed', 'all'}:
        view = 'due'
    q = request.GET.get('q', '').strip()
    month = request.GET.get('month', '').strip()
    week = request.GET.get('week', '').strip()
    stream = request.GET.get('stream', '').strip()
    vendor = request.GET.get('vendor', '').strip()
    scoped, months, month, weeks, week, week_label, stream_names, stream, vendor_names, vendor = _apply_period(
        base, month, week, stream, vendor,
    )
    counts, _streams, _vendors, overdue_q, due_q, upcoming_q, done_q = _pm_analytics(scoped, today, soon)
    tasks = scoped
    if view == 'overdue':
        tasks = tasks.filter(overdue_q).order_by('plan_date', 'equipment_name')
    elif view == 'due':
        tasks = tasks.filter(due_q).order_by('plan_date', 'equipment_name')
    elif view == 'reminder':
        tasks = tasks.filter(completed_on__isnull=True, plan_date__lte=soon).order_by('-plan_date', 'equipment_name')
    elif view == 'upcoming':
        tasks = tasks.filter(upcoming_q).order_by('plan_date', 'equipment_name')
    elif view == 'completed':
        tasks = tasks.filter(done_q).order_by('-completed_on', 'equipment_name')
    if q:
        tasks = tasks.filter(
            Q(value_stream__icontains=q)
            | Q(equipment_name__icontains=q)
            | Q(location__icontains=q)
            | Q(asset_name__icontains=q)
            | Q(asset_number__icontains=q)
        )
    if view == 'all':
        tasks = tasks.order_by('plan_date', 'equipment_name')
    page = Paginator(tasks, 30).get_page(request.GET.get('page'))
    return render(request, 'pm/list.html', {
        'page': page,
        'tasks': page.object_list,
        'counts': counts,
        'view': view,
        'q': q,
        'month': month,
        'months': months,
        'week': week,
        'weeks': weeks,
        'week_label': week_label,
        'stream': stream,
        'stream_names': stream_names,
        'vendor': vendor,
        'vendor_names': vendor_names,
        'today': today,
    })


@roles_required(User.Role.SUPERADMIN, User.Role.ADMIN)
def pm_analytics(request):
    base, today, soon = _pm_base()
    month = request.GET.get('month', '').strip()
    week = request.GET.get('week', '').strip()
    stream = request.GET.get('stream', '').strip()
    vendor = request.GET.get('vendor', '').strip()
    scoped, months, month, weeks, week, week_label, stream_names, stream, vendor_names, vendor = _apply_period(
        base, month, week, stream, vendor,
    )
    counts, streams, vendors, *_rest = _pm_analytics(scoped, today, soon)
    return render(request, 'pm/analytics.html', {
        'counts': counts,
        'streams': streams,
        'vendors': vendors,
        'months': months,
        'month': month,
        'weeks': weeks,
        'week': week,
        'week_label': week_label,
        'stream': stream,
        'stream_names': stream_names,
        'vendor': vendor,
        'vendor_names': vendor_names,
        'today': today,
    })


@roles_required(User.Role.SUPERADMIN, User.Role.ADMIN)
def pm_upload(request):
    upload = request.FILES.get('file')
    if request.method != 'POST' or not upload:
        messages.error(request, 'Choose the preventive maintenance Excel file.')
        return redirect('pm_list')
    if not upload.name.lower().endswith('.xlsx'):
        messages.error(request, 'Upload a .xlsx file.')
        return redirect('pm_list')
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(upload, read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
        result = import_pm_rows(rows)
    except (ValueError, InvalidFileException) as exc:
        messages.error(request, str(exc) or 'That file could not be read.')
        return redirect('pm_list')
    messages.success(request, f"Added {result['created']} planned visits. Skipped {result['skipped']} rows already in the list.")
    return redirect('pm_list')


@roles_required(User.Role.SUPERADMIN, User.Role.ADMIN)
def pm_detail(request, pk):
    task = get_object_or_404(PreventiveTask, pk=pk)
    nxt = PreventiveTask.objects.filter(
        asset_number=task.asset_number,
        equipment_name=task.equipment_name,
        location=task.location,
        value_stream=task.value_stream,
        asset_name=task.asset_name,
        plan_date__gt=task.plan_date,
    ).order_by('plan_date').first()
    return render(request, 'pm/detail.html', {
        'task': task,
        'next_task': nxt,
        'today': timezone.localdate().isoformat(),
    })


@roles_required(User.Role.SUPERADMIN, User.Role.ADMIN)
@require_POST
def pm_complete(request, pk):
    task = get_object_or_404(PreventiveTask, pk=pk)
    if task.completed_on:
        messages.warning(request, 'This visit is already completed.')
        return redirect('pm_detail', pk=task.pk)
    raw = (request.POST.get('completed_on') or '').strip()
    try:
        completed_on = datetime.strptime(raw, '%Y-%m-%d').date()
    except ValueError:
        messages.error(request, 'Choose the date the maintenance was completed.')
        return redirect('pm_detail', pk=task.pk)
    if completed_on > timezone.localdate():
        messages.error(request, 'The completed date cannot be in the future.')
        return redirect('pm_detail', pk=task.pk)
    notes = (request.POST.get('completion_notes') or '').strip()
    task.completed_on = completed_on
    task.completion_notes = notes
    task.completed_by = request.user
    task.save(update_fields=['completed_on', 'completion_notes', 'completed_by'])
    same_asset = PreventiveTask.objects.filter(
        asset_number=task.asset_number,
        equipment_name=task.equipment_name,
        location=task.location,
    ).exclude(pk=task.pk)
    later = same_asset.filter(plan_date__gt=task.plan_date).order_by('plan_date').first()
    back_to_list = request.POST.get('next') == 'list'
    if later:
        messages.success(
            request,
            f'Completed on {completed_on:%d %b %Y}. The next visit is already planned for {later.plan_date:%d %b %Y}.',
        )
        if back_to_list:
            return redirect('pm_list')
        return redirect('pm_detail', pk=task.pk)
    months = task.frequency_months or 3
    next_date = add_months(completed_on, months)
    occupied = PreventiveTask.objects.filter(
        asset_number=task.asset_number,
        equipment_name=task.equipment_name,
        location=task.location,
    )
    guard = 0
    while occupied.filter(plan_date=next_date).exists() and guard < 8:
        next_date = add_months(next_date, months)
        guard += 1
    if occupied.filter(plan_date=next_date).exists():
        messages.success(request, f'Completed on {completed_on:%d %b %Y}.')
        if back_to_list:
            return redirect('pm_list')
        return redirect('pm_detail', pk=task.pk)
    nxt = PreventiveTask.objects.create(
        value_stream=task.value_stream,
        equipment_name=task.equipment_name,
        asset_name=task.asset_name,
        asset_number=task.asset_number,
        location=task.location,
        frequency=task.frequency,
        frequency_months=task.frequency_months or 3,
        schedule_month=schedule_month_label(next_date),
        plan_date=next_date,
        vendor_name=task.vendor_name,
        work_type=task.work_type,
    )
    messages.success(
        request,
        f'Completed on {completed_on:%d %b %Y}. Next maintenance is scheduled for {nxt.plan_date:%d %b %Y}.',
    )
    if back_to_list:
        return redirect('pm_list')
    return redirect('pm_detail', pk=nxt.pk)


def error_403(request, exception=None):
    return render(request, 'errors/403.html', status=403)


def error_404(request, exception=None):
    return render(request, 'errors/404.html', status=404)
