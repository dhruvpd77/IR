import calendar
import os
import re
from datetime import datetime

from django.db import transaction
from django.db.models import Max
from django.utils import timezone

from .models import Asset, Category, InventoryItem, PreventiveTask, StockMovement, SubCategory, Ticket, TicketMedia

MEDIA_EXT = {'.jpg', '.jpeg', '.png', '.webp', '.gif', '.mp4', '.webm', '.mov', '.mkv', '.avi'}
MEDIA_MAX = 40 * 1024 * 1024


def next_ticket_no():
    today = timezone.localdate()
    prefix = f'IRM-{today:%Y%m%d}-'
    latest = (
        Ticket.objects.filter(ticket_no__startswith=prefix)
        .order_by('-ticket_no')
        .values_list('ticket_no', flat=True)
        .first()
    )
    seq = 1
    if latest:
        tail = latest.rsplit('-', 1)[-1]
        if tail.isdigit():
            seq = int(tail) + 1
    return f'{prefix}{seq:03d}'


def validate_media(uploaded):
    ext = os.path.splitext(uploaded.name)[1].lower()
    if ext not in MEDIA_EXT:
        return f'{uploaded.name} is not a supported photo or video.'
    if uploaded.size > MEDIA_MAX:
        return f'{uploaded.name} is larger than 40 MB.'
    return None


def apply_stock_change(*, item, quantity_change, reason, user, ticket=None, note=''):
    if quantity_change == 0:
        raise ValueError('Enter a quantity other than zero.')
    with transaction.atomic():
        locked = InventoryItem.objects.select_for_update().get(pk=item.pk)
        new_qty = locked.quantity + quantity_change
        if new_qty < 0:
            raise ValueError(f'Only {locked.quantity} {locked.unit} of {locked.name} are on hand.')
        locked.quantity = new_qty
        locked.save(update_fields=['quantity'])
        return StockMovement.objects.create(
            item=locked,
            ticket=ticket,
            quantity_change=quantity_change,
            balance_after=new_qty,
            reason=reason,
            note=note,
            created_by=user,
        )


def _make_code(name, existing):
    base = re.sub(r'[^A-Z0-9]+', '-', name.upper()).strip('-')[:12] or 'ITEM'
    code = base
    number = 2
    while code in existing:
        code = f'{base[:8]}-{number}'
        number += 1
    existing.add(code)
    return code


def import_catalog(header, data_rows):
    header = [str(cell or '').strip().lower() for cell in header]

    def find(pred):
        for index, name in enumerate(header):
            if pred(name):
                return index
        return None

    sub_idx = find(lambda name: 'sub' in name and 'desc' not in name)
    cat_idx = find(lambda name: 'category' in name and 'sub' not in name and 'desc' not in name)
    if cat_idx is None or sub_idx is None:
        raise ValueError('The file needs a Category column and a Subcategory column.')
    cat_desc_idx = find(lambda name: 'desc' in name and 'category' in name and 'sub' not in name)
    sub_desc_idx = find(lambda name: 'desc' in name and (name == 'description' or 'sub' in name))

    created_cats = 0
    created_subs = 0
    skipped = 0
    cat_codes = set(Category.objects.values_list('code', flat=True))
    sub_codes = set(SubCategory.objects.values_list('code', flat=True))

    for index, row in enumerate(data_rows, start=2):
        if index > 5001:
            break
        if not row:
            continue
        row = list(row)

        def cell(idx):
            if idx is None or idx >= len(row) or row[idx] is None:
                return ''
            return str(row[idx]).strip()

        cat_name = cell(cat_idx)
        sub_name = cell(sub_idx)
        if not cat_name or not sub_name:
            skipped += 1
            continue
        category = Category.objects.filter(name__iexact=cat_name).first()
        if category is None:
            sort_order = (Category.objects.aggregate(m=Max('sort_order'))['m'] or 0) + 1
            category = Category.objects.create(
                name=cat_name,
                code=_make_code(cat_name, cat_codes),
                description=cell(cat_desc_idx),
                sort_order=sort_order,
            )
            created_cats += 1
        elif cell(cat_desc_idx) and not category.description:
            category.description = cell(cat_desc_idx)
            category.save(update_fields=['description'])
        if SubCategory.objects.filter(category=category, name__iexact=sub_name).exists():
            skipped += 1
            continue
        sub_sort = (category.subcategories.aggregate(m=Max('sort_order'))['m'] or 0) + 1
        SubCategory.objects.create(
            category=category,
            name=sub_name,
            code=_make_code(sub_name, sub_codes),
            description=cell(sub_desc_idx),
            sort_order=sub_sort,
        )
        created_subs += 1
    return {'categories': created_cats, 'subcategories': created_subs, 'skipped': skipped}


def _cell(value):
    if value is None:
        return ''
    return str(value).strip()


def _cmms_triples(rows):
    """Read ValueStream, Equipment Name, and Location, skipping the title row."""
    header_found = False
    indexes = None
    triples = []
    seen = set()
    for row in rows:
        if not row:
            continue
        labels = [_cell(cell).lower() for cell in row]
        if not header_found:
            if 'equipment name' in labels and any('valuestream' in label or label == 'value stream' for label in labels):
                header_found = True
                indexes = {
                    'equipment': labels.index('equipment name'),
                    'stream': next(i for i, label in enumerate(labels) if 'valuestream' in label or label == 'value stream'),
                    'location': labels.index('location'),
                }
            continue
        equipment = _cell(row[indexes['equipment']]) if indexes['equipment'] < len(row) else ''
        stream = _cell(row[indexes['stream']]) if indexes['stream'] < len(row) else ''
        location = _cell(row[indexes['location']]) if indexes['location'] < len(row) else ''
        if not equipment or not stream or not location:
            continue
        key = (stream.casefold(), equipment.casefold(), location.casefold())
        if key in seen:
            continue
        seen.add(key)
        triples.append((stream, equipment, location))
    if not header_found:
        raise ValueError('The file needs columns named ValueStream, Equipment Name, and Location.')
    return triples


def clear_plant_data():
    """Remove requests, stock, and the catalog. User accounts stay."""
    for media in TicketMedia.objects.all():
        if media.file:
            media.file.delete(save=False)
    TicketMedia.objects.all().delete()
    StockMovement.objects.all().delete()
    Ticket.objects.all().delete()
    InventoryItem.objects.all().delete()
    Asset.objects.all().delete()
    SubCategory.objects.all().delete()
    Category.objects.all().delete()


def load_cmms_sheet(path, replace=False):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        triples = _cmms_triples(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    if replace:
        clear_plant_data()
    return import_cmms_triples(triples)


def import_cmms_triples(triples):
    created_streams = 0
    created_equipment = 0
    created_locations = 0
    skipped = 0
    stream_codes = set(Category.objects.values_list('code', flat=True))
    equipment_codes = set(SubCategory.objects.values_list('code', flat=True))
    location_seq = Asset.objects.count()
    with transaction.atomic():
        for stream_name, equipment_name, location_name in triples:
            stream = Category.objects.filter(name__iexact=stream_name).first()
            if stream is None:
                sort_order = (Category.objects.aggregate(m=Max('sort_order'))['m'] or 0) + 1
                stream = Category.objects.create(
                    name=stream_name,
                    code=_make_code(stream_name, stream_codes),
                    sort_order=sort_order,
                )
                created_streams += 1
            equipment = SubCategory.objects.filter(category=stream, name__iexact=equipment_name).first()
            if equipment is None:
                sub_sort = (stream.subcategories.aggregate(m=Max('sort_order'))['m'] or 0) + 1
                equipment = SubCategory.objects.create(
                    category=stream,
                    name=equipment_name,
                    code=_make_code(equipment_name, equipment_codes),
                    sort_order=sub_sort,
                )
                created_equipment += 1
            if Asset.objects.filter(subcategory=equipment, location__iexact=location_name).exists():
                skipped += 1
                continue
            location_seq += 1
            Asset.objects.create(
                subcategory=equipment,
                code=f'EQ-{location_seq:04d}',
                name=equipment_name,
                location=location_name,
            )
            created_locations += 1
    return {
        'categories': created_streams,
        'subcategories': created_equipment,
        'locations': created_locations,
        'skipped': skipped,
    }


def parse_frequency(raw):
    text = str(raw or '').strip().upper()
    if text in {'Y', '1Y', '12M'}:
        return 12
    if text.endswith('M') and text[:-1].isdigit():
        return int(text[:-1])
    return 0


def add_months(value, months):
    month_index = value.month - 1 + months
    year = value.year + month_index // 12
    month = month_index % 12 + 1
    day = min(value.day, calendar.monthrange(year, month)[1])
    return value.replace(year=year, month=month, day=day)


def schedule_month_label(value):
    names = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
    return f"{names[value.month - 1]}'{value.year % 100:02d}"


def _as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if hasattr(value, 'year') and hasattr(value, 'month') and hasattr(value, 'day') and not isinstance(value, str):
        return value
    return None


def import_pm_rows(rows):
    header_index = None
    columns = None
    for index, row in enumerate(rows):
        labels = [_cell(cell).lower() for cell in (row or [])]
        if 'equipment name' in labels and 'location' in labels and any(label.startswith('plan date') for label in labels):
            header_index = index

            def find(*names):
                for name in names:
                    if name in labels:
                        return labels.index(name)
                for offset, label in enumerate(labels):
                    if any(label.startswith(name) for name in names):
                        return offset
                return None

            columns = {
                'equipment': find('equipment name'),
                'stream': find('valuestream', 'value stream'),
                'asset': find('asset name'),
                'number': find('asset number'),
                'location': find('location'),
                'frequency': find('frequency'),
                'month': find('month'),
                'plan': find('plan date'),
                'done': find('done date'),
                'vendor': find('vedor name', 'vendor name'),
                'work': find('external or internal'),
            }
            break
    if columns is None or columns['equipment'] is None or columns['plan'] is None:
        raise ValueError('The file needs Equipment Name, ValueStream, Location, Frequency, Month, and Plan Date.')

    existing = set(
        PreventiveTask.objects.values_list('asset_number', 'equipment_name', 'location', 'plan_date')
    )
    created = 0
    skipped = 0
    batch = []
    for row in rows[header_index + 1:]:
        if not row:
            continue

        def cell(key):
            idx = columns[key]
            if idx is None or idx >= len(row):
                return ''
            return row[idx]

        equipment = _cell(cell('equipment'))
        plan = _as_date(cell('plan'))
        if not equipment or plan is None:
            skipped += 1
            continue
        number = _cell(cell('number'))
        location = _cell(cell('location')) or '—'
        key = (number, equipment, location, plan)
        if key in existing:
            skipped += 1
            continue
        existing.add(key)
        frequency = _cell(cell('frequency')) or '3M'
        months = parse_frequency(frequency) or 3
        done = _as_date(cell('done'))
        work_type = _cell(cell('work'))
        vendor_name = _cell(cell('vendor')) if work_type.lower().startswith('external') else ''
        month_value = cell('month')
        if isinstance(month_value, datetime):
            month_text = schedule_month_label(month_value.date())
        else:
            month_text = _cell(month_value)
            if not month_text or month_text[:4].isdigit():
                month_text = schedule_month_label(plan)
        batch.append(PreventiveTask(
            value_stream=_cell(cell('stream')) or '—',
            equipment_name=equipment,
            asset_name=_cell(cell('asset')),
            asset_number=number,
            location=location,
            frequency=frequency,
            frequency_months=months,
            schedule_month=month_text,
            plan_date=plan,
            completed_on=done,
            vendor_name=vendor_name,
            work_type=work_type,
        ))
        if len(batch) >= 400:
            PreventiveTask.objects.bulk_create(batch)
            created += len(batch)
            batch = []
    if batch:
        PreventiveTask.objects.bulk_create(batch)
        created += len(batch)
    return {'created': created, 'skipped': skipped}
